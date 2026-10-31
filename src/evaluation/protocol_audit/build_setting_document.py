"""Render the verified dev snapshot as an illustrated Word document and HTML.

This is a CPU-only presentation tool. It never edits scores, reruns models, or freezes
an input configuration. Every displayed example resolves through CASE_EVIDENCE.
"""

from __future__ import annotations

import base64
from datetime import datetime, timezone
import hashlib
import html
import json
from pathlib import Path
import sys
import zipfile

from PIL import Image, ImageDraw, ImageFont, ImageOps

sys.path.insert(0, "/tmp/ifbench/doc-runtime")
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path("/tmp/ifbench/tiered_v1/dev_selection")
SNAPSHOT = ROOT / "judge_rejudge13_v1"
OUT = ROOT / "setting_document"
SHARED = ROOT.parent / "shared_inputs_20261027/requests"
NAMES = {
    "qwen2511": "Qwen-Image-Edit-2511",
    "qwen21": "Qwen-Image-2.1",
    "flux3": "FLUX.3 Klein",
    "replan": "RePlan（Qwen-2511）",
    "mirage": "MIRAGE（Qwen-2511）",
    "mice": "MICE（FLUX.2 klein-4B）",
}
DESCRIPTIONS = {
    "S1": "所有目标都有明确文字指代（ref），不提供区域。原生规划器可以从同一公开图文自行预测区域；预测不等于提供 GT 区域。",
    "S2-point": "每个目标均有明确文字指代，同时提供一个点。",
    "S2-box": "每个目标均有明确文字指代，同时提供一个框。",
    "S2-mask": "每个目标均有明确文字指代，同时提供选择器 mask。",
    "S3-point": "指令不含目标的明确文字指代（no-ref），仅由点绑定目标。",
    "S3-box": "指令不含目标的明确文字指代，仅由框绑定目标。",
    "S3-mask": "指令不含目标的明确文字指代，仅由 mask 绑定目标。",
    "S4-ref": "所有目标均有明确文字指代；同一请求内混用 point、box、mask，分配对所有模型固定共享。",
    "S4-noref": "所有目标均无明确文字指代；区域混合分配与 S4-ref 完全相同。",
    "S5": "沿用数据 release 中逐目标的 ref/no-ref 与 none/point/box/mask 组合；none 必须有 ref，不重新随机分配。",
}
COLORS = ["#e33d35", "#2478db", "#2c9c49", "#d29a16", "#b542b5"]


def read(p):
    return json.loads(Path(p).read_text())


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def pct(x):
    return "N/A" if x is None else f"{100 * x:.2f}%"


def score(row, metric):
    lo, hi = row["bootstrap"][metric]["CI95"]
    return f"{pct(row[metric])} [{pct(lo)}, {pct(hi)}]"


def compact(value):
    return json.dumps(value, ensure_ascii=False, separators=(", ", ": "))


def face(size):
    p = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    return ImageFont.truetype(str(p), size) if p.exists() else ImageFont.load_default()


def fit_image(path):
    with Image.open(path) as img:
        return ImageOps.exif_transpose(img).convert("RGB")


def panel(items, path, columns=3):
    """Presentation thumbnails only; never used as model inputs or judged outputs."""
    cols = min(columns, len(items))
    w, h, margin = 600, 450, 18
    rows = (len(items) + cols - 1) // cols
    canvas = Image.new(
        "RGB", (cols * w + (cols + 1) * margin, rows * h + (rows + 1) * margin), "#edf1f5"
    )
    draw = ImageDraw.Draw(canvas)
    for i, (label, img) in enumerate(items):
        left, top = margin + (i % cols) * (w + margin), margin + (i // cols) * (h + margin)
        draw.rectangle((left, top, left + w, top + h), fill="white")
        draw.text((left + 12, top + 9), label, font=face(21), fill="#1b3044")
        preview = img.copy()
        preview.thumbnail((w - 16, h - 54), Image.Resampling.LANCZOS)
        canvas.paste(
            preview, (left + (w - preview.width) // 2, top + 46 + (h - 54 - preview.height) // 2)
        )
    canvas.save(path, quality=91, subsampling=0)
    return path


def rectangle_preview(source, rects):
    img = source.copy()
    draw = ImageDraw.Draw(img)
    width = max(3, round(min(img.size) * 0.004))
    for i, (label, rect) in enumerate(rects):
        x1, y1, x2, y2 = rect
        draw.rectangle(
            (x1, y1, max(x1, x2 - 1), max(y1, y2 - 1)), outline=COLORS[i % 5], width=width
        )
        draw.text(
            (max(0, x1), max(0, y1 - 32)),
            label,
            font=face(30),
            fill=COLORS[i % 5],
            stroke_width=2,
            stroke_fill="white",
        )
    return img


class Report:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        sec.top_margin = Inches(0.65)
        sec.bottom_margin = Inches(0.65)
        sec.left_margin = Inches(0.62)
        sec.right_margin = Inches(0.62)
        for name in ["Normal", "Body Text", "Caption", "Heading 1", "Heading 2", "Heading 3"]:
            style = self.doc.styles[name]
            style.font.name = "Calibri"
            style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        self.doc.styles["Normal"].font.size = Pt(10)
        self.doc.styles["Normal"].paragraph_format.space_after = Pt(5)
        for name in ["Heading 1", "Heading 2", "Heading 3"]:
            self.doc.styles[name].font.color.rgb = RGBColor.from_string("153E5E")
        sec.header.paragraphs[0].text = "IF-Bench  ·  dev 逐 setting 结果与实际输入图册"
        foot = sec.footer.paragraphs[0]
        foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        foot.add_run("dev 审阅版 · 尚未冻结 / 未运行 test    |    ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        foot._p.append(field)
        self.html = [
            '<!doctype html><html lang="zh"><head><meta charset="utf-8"><title>IF-Bench dev逐setting结果与输入图册</title>',
            "<style>body{font:16px/1.65 system-ui,sans-serif;color:#1c3042;max-width:1250px;margin:35px auto;padding:0 24px}h1,h2,h3{color:#153e5e}h2{margin-top:48px;border-bottom:2px solid #dce5ee}table{border-collapse:collapse;width:100%;margin:14px 0;font-size:14px}td,th{border:1px solid #cbd5df;padding:8px;vertical-align:top}th{background:#eaf1f7}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f3f6f9;padding:14px;font-size:13px}img{width:100%;height:auto}figure{margin:20px 0}figcaption{font-size:14px;color:#43566b}p{overflow-wrap:anywhere}.note{background:#fff5de;padding:14px}a{color:#146a9d}@media print{h2{break-before:page}figure{break-inside:avoid}}</style></head><body>",
        ]
        self.images = 0

    def heading(self, text, level=1, anchor=None):
        self.doc.add_heading(text, level=level)
        self.html.append(
            f"<h{level + 1}"
            + (f' id="{anchor}"' if anchor else "")
            + ">"
            + html.escape(text)
            + f"</h{level + 1}>"
        )

    def paragraph(self, text):
        self.doc.add_paragraph(text)
        self.html.append("<p>" + html.escape(text) + "</p>")

    def table(self, heads, rows):
        table = self.doc.add_table(rows=1, cols=len(heads))
        table.style = "Light Shading Accent 1"
        for cell, text in zip(table.rows[0].cells, heads):
            cell.text = str(text)
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for data in rows:
            cells = table.add_row().cells
            for cell, text in zip(cells, data):
                cell.text = str(text)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8 if len(heads) > 5 else 9)
        self.html.append(
            "<table><thead><tr>"
            + "".join("<th>" + html.escape(str(x)) + "</th>" for x in heads)
            + "</tr></thead><tbody>"
        )
        for row in rows:
            self.html.append(
                "<tr>"
                + "".join("<td>" + html.escape(str(x)).replace("\n", "<br>") + "</td>" for x in row)
                + "</tr>"
            )
        self.html.append("</tbody></table>")

    def code(self, label, text):
        self.doc.add_paragraph(label, style="Caption")
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        self.html.append(
            "<p><b>" + html.escape(label) + "</b></p><pre>" + html.escape(text) + "</pre>"
        )

    def picture(self, path, caption):
        with Image.open(path) as im:
            width = im.width
        self.doc.add_picture(str(path), width=Inches(6.95 * min(1, width / 1100)))
        self.doc.add_paragraph(caption, style="Caption")
        self.html.append(
            '<figure><img style="max-width:'
            + str(width)
            + 'px" alt="'
            + html.escape(caption, quote=True)
            + '" src="data:image/jpeg;base64,'
            + base64.b64encode(Path(path).read_bytes()).decode()
            + '"><figcaption>'
            + html.escape(caption)
            + "</figcaption></figure>"
        )
        self.images += 1

    def save(self):
        self.doc.save(OUT / "IF-Bench_dev_setting_report.docx")
        (OUT / "IF-Bench_dev_setting_report.html").write_text("".join(self.html) + "</body></html>")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "figures").mkdir(exist_ok=True)
    analysis_path = SNAPSHOT / "analysis/ANALYSIS.json"
    analysis = read(analysis_path)
    proposal = read(SNAPSHOT / "analysis/PROPOSED_SELECTION.json")
    plan = read(ROOT / "plan.json")
    semantics_path = OUT / "input_semantics_review.json"
    semantics = read(semantics_path)
    input_map = {(x["model"], x["setting"]): x for x in semantics["rows"]}
    assert len(input_map) == 60
    selected = {m: v["candidate"] for m, v in proposal["primary_recommendations"].items()}
    models = list(selected)
    settings = plan["settings"]
    rows = {(v["model"], v["candidate"], v["setting"]): v for v in analysis["rows"]}
    evidence = {
        (v["model"], v["candidate"], v["setting"]): v["cases"]
        for v in read(SNAPSHOT / "analysis/CASE_EVIDENCE.json")["rows"]
    }
    assert (
        analysis["all_rows_score_complete"]
        and not proposal["frozen"]
        and not proposal["test_started"]
    )
    hashes = {
        str(p): sha(p)
        for p in [
            analysis_path,
            SNAPSHOT / "analysis/CASE_EVIDENCE.json",
            SNAPSHOT / "analysis/PROPOSED_SELECTION.json",
            semantics_path,
        ]
    }
    shared_case = plan["case_ids"][0]
    extra_case = next(
        cid
        for cid in plan["case_ids"]
        if len(read(SHARED / "S1" / cid / "public.json")["units"]) == 5
    )
    assert any(
        u["locator"] == "none" for u in read(SHARED / "S5" / extra_case / "public.json")["units"]
    )
    manifest = []
    stat_manifest = []
    report = Report()
    report.doc.add_heading("IF-Bench dev 逐 setting 结果与实际输入图册", 0)
    report.html.append("<h1>IF-Bench dev 逐 setting 结果与实际输入图册</h1>")
    report.paragraph(
        "数据范围：固定 dev 108 张源图、386 个编辑单元；六个模型、10 个 setting。本文展示完整 dev 统计及拟冻结主配置，不是 test 结果，不代表配置已经冻结。生成于 "
        + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        + "。"
    )
    report.paragraph(
        "阅读顺序：先看总览，再按 setting 查看六模型得分、真实输入组成、示例图片和最终编辑 prompt。附录保留全部 180 个候选×setting 统计，以及 S2 最佳适配与 PE/PUP 补充结果。MaskFlow 按用户要求暂不评测；可选及附录模型尚未执行，不能记为 0 分。"
    )
    report.paragraph(
        "UnitPass（UP）：单元 E=4、P≥3、Q≥3 才通过，分母为 386 个单元。JointPass（JP）：一张图全部单元通过且 global P/Q≥3 才通过，分母为 108 张图。所有数值从已审计的 ANALYSIS.json 读取；95% CI 为 10000 次以源图为 cluster 的 paired bootstrap，seed=20261027。百分比按两位小数展示，未调整原始分数。"
    )
    report.paragraph(
        "评分覆盖：11179 个唯一 Judge key 均已有有效评分。首轮 196 条格式失败中，183 条由原答案确定性恢复；剩余 13 条经授权各补评一次后获得有效评分。既有 3 次生成/规划失败仍计入分母。本文只读已有结果，没有增加模型调用。"
    )
    report.paragraph(
        "样例选择不看好坏：每个 setting、每个模型都使用 plan 中第一张源图 "
        + shared_case
        + "（K=2）；另在 S5 增加第一张 K=5 源图 "
        + extra_case
        + "，展示缺区域目标。示例得分与该 setting 的全 dev 统计分别标注。所有示例都插入文档内部；缩略图仅用于展示。"
    )
    report.paragraph(
        "图片约定：Actual input 是接口实际接收的图像内容；Source context 只供读者对照，并不表示单图配置额外接收了原图。Native boxes (report only) 是根据已记录框绘制的文档示意，未送入模型。Binary masks 是 MICE 实际接收的逐目标 mask 缩略图，白色表示所选区域；图像和 mask 的模型内缩放仍按已记录配置执行。模型输出展示统一 resize 回源图后的既有评分图片。"
    )
    report.heading("一、配置与分数总览", 1, "overview")
    report.table(
        ["模型", "拟冻结主配置", "10-setting 平均 JP / 95% CI"],
        [
            [
                NAMES[m],
                selected[m],
                next(
                    pct(c["mean_JointPass"])
                    + " ["
                    + ", ".join(pct(x) for x in c["mean_JointPass_bootstrap"]["CI95"])
                    + "]"
                    for c in proposal["primary_recommendations"][m]["candidate_statistics"]
                    if c["candidate"] == selected[m]
                ),
            ]
            for m in models
        ],
    )
    report.table(
        ["Setting"] + [NAMES[m] for m in models],
        [
            [s]
            + [
                pct(rows[m, selected[m], s]["UnitPass"])
                + " / "
                + pct(rows[m, selected[m], s]["JointPass"])
                for m in models
            ]
            for s in settings
        ],
    )
    report.paragraph(
        "上表每格为 UP / JP。下一节逐行给出对应置信区间。单图 single 配置在 S1 使用原图、在带区域 setting 使用一张标注图；FLUX 的 native-annotation 在有区域时使用原图与标注参考图，并同时提交原生框结构。PE/PUP 在主配置中关闭；原生 planner 按已批准方法流程保留。"
    )
    report.html.append(
        "<nav>" + " · ".join('<a href="#' + s + '">' + s + "</a>" for s in settings) + "</nav>"
    )

    def example(model, setting, cid, extra=False):
        candidate = selected[model]
        ce = next(v for v in evidence[model, candidate, setting] if v["case_id"] == cid)
        directory = ROOT / "runs" / ce["physical_key"]
        call_path = directory / "call.json"
        result_path = directory / "result.json"
        call, result = read(call_path), read(result_path)
        for p in [call_path, result_path]:
            hashes[str(p)] = sha(p)
        source_path = Path(call["source"]["path"])
        assert sha(source_path) == call["source"]["sha256"]
        source = fit_image(source_path)
        title = NAMES[model] + (" · S5 缺区域补例" if extra else "")
        report.heading(title, 2)
        report.paragraph(input_map[model, setting]["description"])
        for field, label in [
            ("coordinates", "坐标与转换"),
            ("default_components", "默认组件"),
            ("notes", "补充说明"),
        ]:
            if input_map[model, setting].get(field):
                value = input_map[model, setting][field]
                report.paragraph(
                    label + "：" + (value if isinstance(value, str) else compact(value))
                )
        report.paragraph(
            f"样例 {cid}，K={ce['K']}；实际配置 {candidate}；该例 UP={ce['UnitPass_count']}/{ce['K']}，JP={ce['JointPass']}；生成状态={ce['generation_status']}。这是一例的通过情况，不是全 setting 通过率。"
        )
        units = read(SHARED / setting / cid / "public.json")["units"]
        report.paragraph(
            "公开目标与区域："
            + "；".join(
                u["id"] + " = " + ("ref" if u["has_ref"] else "no-ref") + " / " + u["locator"]
                for u in units
            )
            + "。"
        )
        items = []
        inputs = [Path(p) for p in call["images"]]
        source_inputs = list(inputs)
        if model == "mice":
            actual = directory / "actual_source.png"
            if not actual.exists() and result.get("reuse_record"):
                actual = Path(result["reuse_record"]).parent / "actual_source.png"
            assert actual.exists(), "MICE actual source image missing"
            assert list(fit_image(actual).size) == call["model_size"]
            inputs = [actual]
        if source_path not in inputs:
            items.append(("Source context (not extra input)", source))
        for i, p in enumerate(inputs):
            items.append((f"Actual input {i + 1}", fit_image(p)))
            hashes[str(p)] = sha(p)
        if result.get("output_image"):
            output = Path(result["output_image"])
            assert sha(output) == result["output_sha256"]
            hashes[str(output)] = sha(output)
            items.append(("Actual output (source size)", fit_image(output)))
        key = f"{setting}__{model}__{cid}"
        figure = panel(items, OUT / "figures" / f"{key}.jpg")
        report.picture(
            figure, f"{setting} / {NAMES[model]} / {cid}：实际输入与既有输出；按原比例缩略。"
        )
        rects = []
        geometry = []
        if model == "replan":
            geometry = result.get("actual_region_guidance", call.get("regions", []))
            rects = [(f"R{i + 1}", r.get("bbox", r.get("bbox_2d"))) for i, r in enumerate(geometry)]
        elif model == "mirage":
            geometry = call.get("regions", [])
            rects = [
                (r.get("unit_id", f"R{i + 1}"), [r["bbox"][k] for k in ["x1", "y1", "x2", "y2"]])
                for i, r in enumerate(geometry)
            ]
        elif model == "flux3":
            geometry = call.get("boxes", [])
        if rects:
            p = panel(
                [("Native boxes (report only)", rectangle_preview(source, rects))],
                OUT / "figures" / f"{key}__regions.jpg",
            )
            report.picture(
                p,
                "原生区域框的文档示意，未作为额外标记图输入。RePlan 展示执行记录中的实际框；MIRAGE 展示传入的区域框，后续仍按官方 crop/context 流程处理。",
            )
        masks = call.get("masks", [])
        if masks:
            ids = call.get("regional_unit_ids", [])
            assert len(ids) == len(masks)
            p = panel(
                [(f"Binary mask {uid}", fit_image(path)) for uid, path in zip(ids, masks)],
                OUT / "figures" / f"{key}__masks.jpg",
                columns=3,
            )
            report.picture(
                p,
                "MICE 实际 mask 输入；白色为选择区域，黑色为其余区域。mask 与后面的区域指令按 regional_unit_ids 顺序一一对应。",
            )
            for p in masks:
                hashes[str(p)] = sha(p)
        if geometry:
            if model == "replan":
                report.table(
                    ["区域", "实际像素 xyxy 框", "局部指令"],
                    [
                        [f"R{i + 1}", compact(r.get("bbox", r.get("bbox_2d"))), r["hint"]]
                        for i, r in enumerate(geometry)
                    ],
                )
            elif model == "mirage":
                report.table(
                    ["目标", "传入像素 xyxy 框", "refer_object"],
                    [
                        [
                            r.get("unit_id", f"R{i + 1}"),
                            compact([r["bbox"][k] for k in ["x1", "y1", "x2", "y2"]]),
                            r["refer_object"],
                        ]
                        for i, r in enumerate(geometry)
                    ],
                )
            elif model == "flux3":
                report.code(
                    "原生 boxes 入参（归一化坐标，约定详见输入方式）",
                    json.dumps(geometry, ensure_ascii=False, indent=2),
                )
        final_prompt = call.get("prompt", "")
        prompt_origin = "call.json: prompt"
        if model == "replan":
            final_prompt = result["actual_global_prompt"]
            prompt_origin = "result.json: actual_global_prompt；局部 hint 见上表"
        elif model == "flux3":
            final_prompt = result.get("generation_metadata", {}).get("prompt")
            assert isinstance(final_prompt, str), "FLUX final serialized prompt missing"
            prompt_origin = "result.json: generation_metadata.prompt（含原生 boxes 序列化）"
        report.code("实际最终编辑 prompt · " + prompt_origin, final_prompt)
        if model == "mirage":
            report.code(
                "实际区域 sub_prompts（原顺序）",
                "\n\n".join(f"{i + 1}. {p}" for i, p in enumerate(call["sub_prompts"])),
            )
        if model == "mice":
            report.paragraph(
                "文本槽位：global_unit_ids="
                + compact(call.get("global_unit_ids", []))
                + "；regional_unit_ids="
                + compact(call.get("regional_unit_ids", []))
                + "。字面量 $BREAKFLAG$ 分隔全局槽与区域槽，不替换为普通换行。模型输入尺寸="
                + compact(call["model_size"])
                + "。"
            )
        report.paragraph(
            "该请求记录的附加组件调用："
            + compact(call.get("extra_component_calls", {}))
            + "。这些是请求的逻辑流程记录；复用缓存的调用不应重复加总为新增物理调用。"
        )
        params = call.get("parameters") or result.get("configuration")
        if params:
            report.code("保存的推理参数", compact(params))
        report.paragraph(
            "证据："
            + str(call_path)
            + "；"
            + str(result_path)
            + "。"
            + (
                "此逻辑请求复用 " + ce["physical_key"] + " 的同一结果。"
                if ce["physical_key"] != ce["logical_key"]
                else ""
            )
        )
        item = {
            "model": model,
            "candidate": candidate,
            "setting": setting,
            "case_id": cid,
            "additional_S5_example": extra,
            "logical_key": ce["logical_key"],
            "physical_key": ce["physical_key"],
            "call": str(call_path),
            "result": str(result_path),
            "input_images": [str(x) for x in inputs],
            "call_source_images": [str(x) for x in source_inputs],
            "mask_images": masks,
            "output_image": result.get("output_image"),
            "figure": str(figure),
            "final_prompt": final_prompt,
            "final_prompt_origin": prompt_origin,
            "UnitPass_count": ce["UnitPass_count"],
            "K": ce["K"],
            "JointPass": ce["JointPass"],
            "judge_key": ce["judge_key"],
        }
        manifest.append(item)

    for setting in settings:
        report.heading("二、" + setting, 1, setting)
        report.paragraph(DESCRIPTIONS[setting])
        statistics = []
        for m in models:
            row = rows[m, selected[m], setting]
            statistics.append(
                [
                    NAMES[m],
                    selected[m],
                    score(row, "UnitPass"),
                    score(row, "JointPass"),
                    row["generation_failed_cases"],
                ]
            )
            stat_manifest.append(
                {
                    k: row[k]
                    for k in [
                        "model",
                        "candidate",
                        "setting",
                        "UnitPass",
                        "JointPass",
                        "bootstrap",
                        "case_denominator",
                        "unit_denominator",
                        "generation_failed_cases",
                    ]
                }
            )
        report.table(
            ["模型", "输入配置", "UP [95% CI]", "JP [95% CI]", "生成/规划失败图数"], statistics
        )
        report.paragraph(
            "全部行使用 108 图 / 386 单元完整分母；上表失败数是候选×setting 的逻辑计数，不能跨别名行相加当作物理失败次数。"
        )
        report.table(
            ["模型", "该 setting 下实际怎么输入"],
            [[NAMES[m], input_map[m, setting]["description"]] for m in models],
        )
        for m in models:
            example(m, setting, shared_case)
        if setting == "S5":
            report.heading("S5 补例：同一请求中有缺区域目标（K=5）", 2)
            report.paragraph(
                "固定首个 K=5 源图中 U2 没有区域但有明确文字指代。这里展示 RePlan/MIRAGE 原生规划、MICE 全局槽，以及通用编辑模型的文字路径；不把未提供的区域补成 GT。"
            )
            for m in models:
                example(m, setting, extra_case, extra=True)

    report.heading("三、S2 最佳适配补充结果", 1, "s2-supplement")
    report.paragraph(
        "下表是单独的 S2 best-adapted 补充，不能替换上文保留区域的主配对行。discard-region 只在 S2 使用 S1 输入；绝不表示 S3/S4-noref 丢弃区域。同分先保留区域，再按预登记候选顺序。"
    )
    supp = []
    for m in models:
        for setting, v in proposal["S2_best_adapted"][m].items():
            row = rows[m, v["candidate"], setting]
            supp.append(
                [
                    NAMES[m],
                    setting,
                    v["candidate"],
                    "丢弃区域，复用 S1" if v["discard_region"] else "保留区域",
                    score(row, "UnitPass"),
                    score(row, "JointPass"),
                ]
            )
    report.table(["模型", "setting", "候选", "区域使用", "UP [95% CI]", "JP [95% CI]"], supp)
    report.heading("四、全部候选统计（含官方组件补充）", 1, "all-candidates")
    report.paragraph(
        "共 180 行。primary 是主配置候选，official-supplement 是 PE/PUP 完整系统补充，discard 是 S2 回退候选。discard 候选在非 S2 的行只是首个主候选原请求复用，不表示所有 setting 都不使用区域。PE/PUP 行不参与关闭改写组件的主行选择。"
    )
    report.table(
        ["候选", "输入方式说明"],
        [
            ["single", "S1 原图；有区域时单张带标记图＋绑定颜色/编号的指令。"],
            ["dual", "原图＋区域标记副本，prompt 指定只编辑原图；S1 复用无标记原图请求。"],
            [
                "coordinates",
                "原图＋显式坐标文字；具体坐标序列化和有损转换以已登记 call.json 为准。",
            ],
            [
                "native / native-annotation",
                "原生输入；FLUX native-annotation 额外加入标注参考图。其他模型的原生方式逐 setting 见正文。",
            ],
            [
                "official-pe / official-pup",
                "Qwen-2.1 官方 Prompt Enhancement / FLUX 官方 Prompt Upsampling 开启，单列系统补充。",
            ],
            [
                "discard-region",
                "仅 S2 丢弃区域、复用 S1；其余 setting 为登记首个主候选的原请求复用。",
            ],
        ],
    )
    for m in models:
        report.heading(NAMES[m] + "：全部候选", 2)
        report.table(
            ["候选", "角色", "setting", "UP [95% CI]", "JP [95% CI]"],
            [
                [
                    v["candidate"],
                    v["role"],
                    v["setting"],
                    score(v, "UnitPass"),
                    score(v, "JointPass"),
                ]
                for v in analysis["rows"]
                if v["model"] == m
            ],
        )
    report.heading("五、来源、限制与状态", 1, "provenance")
    report.paragraph(
        "数据/模型输入/编辑输出/评分均为既有快照；本文生成过程只进行读取、文档排版和可视化缩略图制作。文档中的原生框示意仅可由保存的公开/预测框绘制，未读取隐藏 GT 为编辑输入补区域。样例均按固定规则抽取，不估计总体标记残留率。"
    )
    report.paragraph(
        "人工双人标记残留检查、GT selector paste-back 诊断、其余未执行消融以及正式 test 仍属后续工作。冻结前确认节点保持不变。"
    )
    for p in [
        analysis_path,
        SNAPSHOT / "analysis/CASE_EVIDENCE.json",
        SNAPSHOT / "analysis/PROPOSED_SELECTION.json",
        semantics_path,
    ]:
        report.paragraph(str(p) + "\nSHA256 " + hashes[str(p)])
    report.paragraph(
        "可复现生成器："
        + str(Path(__file__).resolve())
        + "。同目录 DOCUMENT_MANIFEST.json 保存 66 组示例、60 个主统计行、输入输出路径与内容 SHA256；全部图像直接嵌入 docx/HTML，无需联网加载。"
    )
    report.save()
    for p, h in hashes.items():
        assert sha(p) == h, "Source artifact changed: " + p
    assert len(manifest) == 66 and len(stat_manifest) == 60
    docx_path = OUT / "IF-Bench_dev_setting_report.docx"
    with zipfile.ZipFile(docx_path) as z:
        assert z.testzip() is None
        xml = z.read("word/document.xml").decode()
        assert xml.count("<w:drawing>") == report.images
        assert not any(
            b'TargetMode="External"' in z.read(p) for p in z.namelist() if p.endswith(".rels")
        )
        embedded_media = len([p for p in z.namelist() if p.startswith("word/media/")])
    validation = {
        "main_model_setting_examples": 60,
        "additional_S5_K5_examples": 6,
        "main_statistics_rows": 60,
        "appendix_candidate_setting_rows": len(analysis["rows"]),
        "embedded_figure_occurrences": report.images,
        "unique_docx_media": embedded_media,
        "input_output_hashes_verified": len(hashes),
        "source_artifacts_unchanged": True,
        "docx_zip_valid": True,
        "external_image_links": 0,
        "new_model_calls": 0,
        "frozen": False,
        "test_started": False,
    }
    (OUT / "DOCUMENT_MANIFEST.json").write_text(
        json.dumps(
            {
                "created_at_utc": datetime.now(timezone.utc).isoformat(),
                "sample_rule": "first plan case for all 60 primary rows; first K5 case as 6 additional S5 examples; no outcome filtering",
                "examples": manifest,
                "statistics": stat_manifest,
                "source_sha256": hashes,
                "validation": validation,
                "artifacts_sha256": {
                    p.name: sha(p) for p in [docx_path, OUT / "IF-Bench_dev_setting_report.html"]
                },
                "generator_sha256": sha(__file__),
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n"
    )
    (OUT / "VALIDATION.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2) + "\n"
    )
    print(json.dumps(validation, ensure_ascii=False))


if __name__ == "__main__":
    main()
